"""Index the two evidence-map Word files, one JSONL + one summary per unit (Ch 4 to the Afterword by default).

Free, local, no API. Citation lines are found by their ID pattern (never by Word style); the quotation of an item is
the run of non-heading paragraphs since the previous citation or heading.

  python -I -X utf8 drafting/_tools/index_evidence_maps.py            # units C04..C13
  python -I -X utf8 drafting/_tools/index_evidence_maps.py --all      # every unit

Writes drafting/evidence_index/chNN.jsonl and chNN_summary.md (ch13 = Afterword), and prints the checks:
item count per unit against each map's chapter table; unique IDs; IDs resolving in drafting/_evidence_log.jsonl.
"""
from __future__ import annotations

import collections
import json
import re
import sys
from pathlib import Path

import docx

HERE = Path(__file__).resolve().parent
DR = HERE.parent
ROOT = DR.parent
MAPS = ROOT / "Udaypur Evidence Mapping Documents"
FILES = {"direct": MAPS / "Udaypur Evidence Map - Direct Evidence.docx",
         "contextual": MAPS / "Udaypur Evidence Map - Contextual Evidence.docx"}
OUT = DR / "evidence_index"
LOG = DR / "_evidence_log.jsonl"
sys.path.insert(0, str(DR / "_pipe"))
import tiers  # noqa: E402

CIT = re.compile(r"^(C\d\d-[A-Z0-9]+(?:-[A-Z0-9]+)*-\d{3})\s{2,}(.*)$")
BOOK = re.compile(r"^(.*)\(([A-Z0-9-]+), (\d+)\)\s*$")
CHAP = re.compile(r"^(Chapter (\d+)|Afterword)\b")


def unit_of(heading: str) -> str | None:
    m = CHAP.match(heading.strip())
    if not m:
        return None
    return "C13" if m.group(1) == "Afterword" else f"C{int(m.group(2)):02d}"


def parse(kind: str, path: Path):
    d = docx.Document(str(path))
    table = {}
    for t in d.tables:
        if t.rows and t.rows[0].cells[0].text.strip() == "Chapter":
            for r in t.rows[1:]:
                u = unit_of(r.cells[0].text)
                if u:
                    table[u] = int(re.sub(r"\D", "", r.cells[2].text) or 0)
    items, unit, book, quote = [], None, None, []
    for p in d.paragraphs:
        txt = p.text.strip()
        style = p.style.name if p.style is not None else ""
        if not txt:
            continue
        m = CIT.match(txt)
        if m:
            rest = m.group(2)
            loc, _, also = rest.partition("also ")
            loc = re.sub(r"\s{2,}", " ", loc).strip()
            also_l = [a.strip() for a in also.split(",") if a.strip()] if also else []
            flags = []
            if "⚠" in txt or "⚠" in " ".join(quote):
                flags.append("⚠")
            if "≈" in loc:
                flags.append("≈")
            if "no page" in loc.lower() or "not recorded" in loc.lower():
                flags.append("no page")
            if also_l:
                flags.append("also " + ", ".join(also_l))
            items.append({"id": m.group(1), "unit": unit, "kind": kind, "book_code": book[0] if book else None,
                          "book_title": book[1] if book else None, "page": loc, "also": also_l, "flags": flags,
                          "quote": "\n".join(quote).strip()})
            quote = []
            continue
        if style.startswith("Heading 1"):
            u = unit_of(txt)
            if u:
                unit, book = u, None
            quote = []
            continue
        if style.startswith("Heading 2"):
            b = BOOK.match(txt)
            book = (b.group(2), b.group(1).strip()) if b else (None, txt)
            quote = []
            continue
        if style.startswith("Heading") or style == "Title":
            quote = []
            continue
        quote.append(txt)
    return items, table


def main():
    want = None if "--all" in sys.argv else {f"C{n:02d}" for n in range(4, 14)}
    log = {}
    if LOG.exists():
        for line in LOG.open(encoding="utf-8"):
            r = json.loads(line)
            log[r["id"]] = r
    all_items, tables = [], {}
    for kind, path in FILES.items():
        it, tb = parse(kind, path)
        all_items += it
        tables[kind] = tb
    # checks
    ids = collections.Counter(i["id"] for i in all_items)
    dups = [k for k, v in ids.items() if v > 1]
    counts = collections.Counter((i["unit"], i["kind"]) for i in all_items)
    report = ["# Evidence index: build checks", ""]
    report.append(f"Items parsed: {len(all_items)} ({sum(1 for i in all_items if i['kind']=='direct')} direct, "
                  f"{sum(1 for i in all_items if i['kind']=='contextual')} contextual). Duplicate IDs: {dups or 'none'}.")
    report += ["", "| Unit | Direct (parsed / map table) | Contextual (parsed / map table) | Match |", "|---|---|---|---|"]
    units = sorted({i["unit"] for i in all_items if i["unit"]})
    for u in units:
        dp, dt = counts[(u, "direct")], tables["direct"].get(u)
        cp, ct = counts[(u, "contextual")], tables["contextual"].get(u)
        report.append(f"| {u} | {dp} / {dt} | {cp} / {ct} | {'yes' if dp == dt and cp == ct else '**NO**'} |")
    unresolved = [i["id"] for i in all_items if i["id"] not in log]
    report += ["", f"IDs not found in `_evidence_log.jsonl`: {len(unresolved)}" + (f" ({', '.join(unresolved[:40])}{' …' if len(unresolved) > 40 else ''})" if unresolved else "")]
    mism = []
    for i in all_items:
        r = log.get(i["id"])
        if r and r.get("book_code") and i["book_code"] and r["book_code"] != i["book_code"]:
            mism.append(f"{i['id']}: map {i['book_code']} / log {r['book_code']}")
    report.append(f"Book-code mismatches map vs log: {len(mism)}" + (f" ({'; '.join(mism[:20])})" if mism else ""))
    OUT.mkdir(exist_ok=True)
    (OUT / "BUILD_CHECKS.md").write_text("\n".join(report) + "\n", encoding="utf-8")
    # per unit
    for u in units:
        if want and u not in want:
            continue
        rows = [i for i in all_items if i["unit"] == u]
        for i in rows:
            i["tier"], i["tier_note"] = tiers.tier_of(i)
            r = log.get(i["id"], {})
            i["sub"], i["strength"], i["claim"], i["gloss"] = r.get("sub"), r.get("strength"), r.get("claim"), r.get("gloss")
            i["in_log"] = bool(r)
        n = u[1:]
        with (OUT / f"ch{n}.jsonl").open("w", encoding="utf-8") as f:
            for i in rows:
                f.write(json.dumps(i, ensure_ascii=False) + "\n")
        s = [f"# {u} evidence index summary", "",
             f"{len(rows)} items: {sum(1 for i in rows if i['kind']=='direct')} direct, {sum(1 for i in rows if i['kind']=='contextual')} contextual. "
             f"Full quotations: `ch{n}.jsonl`.", "", "## By book", "", "| Book | Tier | Direct | Contextual |", "|---|---|---|---|"]
        bb = collections.defaultdict(lambda: [0, 0, None])
        for i in rows:
            b = bb[i["book_code"]]
            b[0 if i["kind"] == "direct" else 1] += 1
            b[2] = i["tier"]
        for k, (dd, cc, t) in sorted(bb.items(), key=lambda x: -(x[1][0] + x[1][1])):
            s.append(f"| {k} | {t} | {dd} | {cc} |")
        s += ["", "## By sub-topic (from the evidence log)", ""]
        for k, v in sorted(collections.Counter(i["sub"] for i in rows).items(), key=lambda x: str(x[0])):
            s.append(f"- {k}: {v}")
        s += ["", "## Also-Ch-N cross-references", ""]
        xc = collections.Counter(a for i in rows for a in i["also"])
        s.append(", ".join(f"{k}: {v}" for k, v in sorted(xc.items())) or "none")
        s += ["", "## Flagged items (⚠, no page)", ""]
        fl = [i for i in rows if "⚠" in i["flags"] or "no page" in i["flags"]]
        s.append(f"{len(fl)} items: " + ", ".join(i["id"] for i in fl) if fl else "none")
        s += ["", "## Compact index", "", "ID | kind | book (tier) | page | sub | first 120 characters", ""]
        for i in rows:
            q = re.sub(r"\s+", " ", i["quote"])[:120]
            s.append(f"- {i['id']} | {i['kind'][0].upper()} | {i['book_code']} (T{i['tier']}) | {i['page']} | {i['sub']} | {q}")
        (OUT / f"ch{n}_summary.md").write_text("\n".join(s) + "\n", encoding="utf-8")
    print("\n".join(report))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
