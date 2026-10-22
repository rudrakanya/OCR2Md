"""Footnoted Markdown -> .docx with real Word footnotes (pandoc 3.9 via pypandoc_binary in drafting/_tools/vendor).

  set UDAY_CFG=...; python -I drafting/_pipe/make_docx.py [footnoted.md] [out.docx]

Layout: A4, 1-inch margins; Georgia 11 pt body, 1.25 line spacing; a small "CHAPTER N" label, the title, then
"Draft vN, with footnotes · October 2026"; section headings Heading 1; footnotes 9 pt, each ending with its evidence
IDs in small grey text; bibliography on its own page, each book once, hanging indent; page-number footer.
"""
import re
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from cfg import C, ROOT  # noqa: E402
sys.path.insert(0, str(ROOT / "drafting" / "_tools" / "vendor"))
import pypandoc  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

PANDOC = pypandoc.get_pandoc_path()


def prepare(md: str) -> str:
    md = re.sub(r"(?m)^---\s*\n+## Notes\s*\n", "\n", md)
    # evidence IDs (HTML comments after each note) -> small grey span
    md = re.sub(r"\s*<!--\s*(C\d\d-[^>]*?)\s*-->", lambda m: " [" + m.group(1).strip() + ']{custom-style="Evidence ID"}', md)
    md = re.sub(r"(?m)^## (\d+\.\d+ .*)$", r"# \1", md)
    if not getattr(C, "NUMBERED_HEADINGS", True):          # the Afterword prints its units unnumbered
        md = re.sub(r"(?m)^# \d+\.\d+ (.*)$", r"# \1", md)
    head, _, bib = md.partition("\n## Bibliography\n")
    pagebreak = '\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n\n'
    title = (f'::: {{custom-style="Chapter Label"}}\n{getattr(C, "CHAPTER_LABEL", f"CHAPTER {C.CH}")}\n:::\n\n::: {{custom-style="Title"}}\n{C.TITLE}\n:::\n\n'
             '::: {custom-style="Subtitle"}\n' + getattr(C, "SUBTITLE", f"Draft {C.VERSION}, with footnotes · October 2026") + '\n:::\n\n')
    out = title + head.rstrip() + "\n"
    if bib:
        entries = [e.strip() for e in bib.strip().split("\n\n") if e.strip()]
        out += pagebreak + "# Bibliography\n\n" + "\n\n".join(f'::: {{custom-style="Bibliography Entry"}}\n{e}\n:::' for e in entries) + "\n"
    return out


def _font(style, name="Georgia", size=None, bold=None, italic=None, color=None):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def _style(doc, name, kind=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, kind)


def finish(path: Path):
    doc = Document(str(path))
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        _font(st, size=11)
        st.paragraph_format.line_spacing = 1.25
    for name in ("Footnote Text", "Footnote Block Text"):
        try:
            st = doc.styles[name]
            _font(st, size=9)
            st.paragraph_format.line_spacing = 1.0
            st.paragraph_format.space_after = Pt(2)
        except KeyError:
            pass
    h1 = doc.styles["Heading 1"]
    _font(h1, size=14, bold=True, color="222222")
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(6)
    _font(_style(doc, "Title"), size=22, bold=False, color="111111")
    _font(_style(doc, "Subtitle"), size=11, italic=True, color="555555")
    lab = _style(doc, "Chapter Label")
    _font(lab, size=9, bold=True, color="777777")
    lab.paragraph_format.space_after = Pt(4)
    rpr = lab.element.get_or_add_rPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), "40")
    rpr.append(sp)
    ev = _style(doc, "Evidence ID", WD_STYLE_TYPE.CHARACTER)
    _font(ev, size=7, color="8C8C8C")
    bib = _style(doc, "Bibliography Entry")
    _font(bib, size=11)
    bib.paragraph_format.left_indent = Inches(0.5)
    bib.paragraph_format.first_line_indent = Inches(-0.5)
    bib.paragraph_format.line_spacing = 1.15
    bib.paragraph_format.space_after = Pt(6)
    for sec in doc.sections:
        sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
        p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        p.alignment = 1
        run = p.add_run()
        for kind, txt in (("begin", None), (None, " PAGE "), ("end", None)):
            if kind:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                run._r.append(fc)
            else:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = txt
                run._r.append(it)
        run.font.size = Pt(9)
        run.font.name = "Georgia"
    doc.save(str(path))


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else C.OUT.with_name(C.OUT.stem + "_footnoted.md")
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else C.D / C.DOCX_NAME
    md = prepare(src.read_text(encoding="utf-8"))
    tmp = C.D / "_docx_input.md"
    tmp.write_text(md, encoding="utf-8")
    subprocess.run([PANDOC, str(tmp), "-f", "markdown+footnotes+raw_attribute+fenced_divs+bracketed_spans", "-t", "docx", "-o", str(dst)], check=True)
    tmp.unlink()
    finish(dst)
    import zipfile
    fx = zipfile.ZipFile(dst).read("word/footnotes.xml").decode("utf-8")
    n = len([i for i in re.findall(r'<w:footnote [^>]*w:id="(-?\d+)"', fx) if int(i) > 0])
    print(f"wrote {dst.name}: {n} Word footnotes; evidence-ID spans in notes: {fx.count('EvidenceID') + fx.count('Evidence ID')}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
